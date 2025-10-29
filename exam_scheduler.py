import tkinter as tk
from tkinter import filedialog, messagebox
from tkinter import ttk
from datetime import date, datetime, timedelta
import json

from tkcalendar import Calendar
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVER_GREEN = "#03C75A"
NAVER_DARK = "#1E1E1E"
NAVER_BG = "#F8F9FA"
NAVER_CARD = "#FFFFFF"


class Teacher:
    def __init__(self, name):
        self.name = name
        self.unavailable = set()
        self.assignments = {"main": 0, "assistant": 0, "selfstudy": 0}

    def add_unavailability(self, date_str, period):
        self.unavailable.add((date_str, period))


class Subject:
    def __init__(self, name, teacher):
        self.name = name
        self.teacher = teacher


class ExamScheduler(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("시험 감독 스케줄러")
        self.geometry("1000x680")
        self.configure(bg=NAVER_BG)

        self.exam_dates = {}
        self.subjects = []
        self.teachers = {}
        self.grade_count = 1
        self.class_count = 1
        self.generated_schedule = []
        self.max_periods = 4

        self.setup_style()
        self.create_widgets()

    def setup_style(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass

        style.configure("TNotebook", background=NAVER_BG, borderwidth=0)
        style.configure(
            "TNotebook.Tab",
            padding=(18, 8),
            font=("Pretendard", 10, "bold"),
        )
        style.map(
            "TNotebook.Tab",
            background=[("selected", NAVER_CARD)],
            foreground=[("selected", NAVER_GREEN)],
        )
        style.configure("Card.TFrame", background=NAVER_CARD)
        style.configure("TFrame", background=NAVER_CARD)
        style.configure(
            "Naver.TLabel",
            background=NAVER_CARD,
            foreground=NAVER_DARK,
            font=("Pretendard", 10),
        )
        style.configure(
            "NaverHeading.TLabel",
            background=NAVER_CARD,
            foreground=NAVER_DARK,
            font=("Pretendard", 11, "bold"),
        )
        style.configure(
            "Naver.TButton",
            font=("Pretendard", 10, "bold"),
            background=NAVER_GREEN,
            foreground="#FFFFFF",
            padding=(18, 10),
            borderwidth=0,
        )
        style.map(
            "Naver.TButton",
            background=[("active", "#02B456"), ("disabled", "#B5EBD1")],
            foreground=[("disabled", "#FFFFFF")],
        )
        style.configure(
            "Secondary.TButton",
            font=("Pretendard", 10),
            background="#E5F8ED",
            foreground=NAVER_GREEN,
            padding=(14, 8),
        )
        style.map(
            "Secondary.TButton",
            background=[("active", "#D1F2E0")],
        )
        style.configure("Naver.TSpinbox", font=("Pretendard", 10))
        style.configure("Naver.TEntry", font=("Pretendard", 10))

    def create_widgets(self):
        header = tk.Frame(self, bg=NAVER_GREEN, height=72)
        header.pack(fill="x")
        tk.Label(
            header,
            text="NE Teacher Exam Scheduler",
            bg=NAVER_GREEN,
            fg="white",
            font=("Pretendard", 18, "bold"),
            pady=12,
        ).pack(side="left", padx=24)

        container = ttk.Frame(self, style="Card.TFrame")
        container.pack(fill="both", expand=True, padx=24, pady=24)

        notebook = ttk.Notebook(container)
        notebook.pack(fill="both", expand=True)

        self.period_frame = ttk.Frame(notebook, style="Card.TFrame", padding=24)
        notebook.add(self.period_frame, text="시험 기간")
        self.subject_frame = ttk.Frame(notebook, style="Card.TFrame", padding=24)
        notebook.add(self.subject_frame, text="과목")
        self.teacher_frame = ttk.Frame(notebook, style="Card.TFrame", padding=24)
        notebook.add(self.teacher_frame, text="교사")
        self.schedule_frame = ttk.Frame(notebook, style="Card.TFrame", padding=24)
        notebook.add(self.schedule_frame, text="스케줄")

        self.build_period_frame()
        self.build_subject_frame()
        self.build_teacher_frame()
        self.build_schedule_frame()

    # ----- 시험 기간 설정 -----
    def build_period_frame(self):
        frm = self.period_frame
        frm.columnconfigure(5, weight=1)
        frm.rowconfigure(4, weight=1)

        ttk.Label(frm, text="시험 기본 설정", style="NaverHeading.TLabel").grid(
            row=0, column=0, columnspan=6, sticky="w", pady=(0, 12)
        )

        ttk.Label(frm, text="학년 수", style="Naver.TLabel").grid(row=1, column=0, sticky="e")
        ttk.Label(frm, text="학급 수", style="Naver.TLabel").grid(row=1, column=2, sticky="e")
        self.grade_spin = ttk.Spinbox(
            frm,
            from_=1,
            to=12,
            width=6,
            style="Naver.TSpinbox",
            justify="center",
        )
        self.class_spin = ttk.Spinbox(
            frm,
            from_=1,
            to=20,
            width=6,
            style="Naver.TSpinbox",
            justify="center",
        )
        self.grade_spin.set("1")
        self.class_spin.set("1")
        self.grade_spin.grid(row=1, column=1, padx=4, pady=4)
        self.class_spin.grid(row=1, column=3, padx=4, pady=4)

        ttk.Label(frm, text="하루 교시 수", style="Naver.TLabel").grid(row=1, column=4, sticky="e")
        self.max_period_spin = ttk.Spinbox(
            frm,
            from_=1,
            to=10,
            width=6,
            style="Naver.TSpinbox",
            justify="center",
        )
        self.max_period_spin.set("4")
        self.max_period_spin.grid(row=1, column=5, padx=4, pady=4, sticky="w")

        calendar_card = ttk.Frame(frm, style="Card.TFrame")
        calendar_card.grid(row=2, column=0, columnspan=6, sticky="nsew", pady=(20, 12))
        calendar_card.columnconfigure(0, weight=1)

        self.cal = Calendar(
            calendar_card,
            selectmode="day",
            year=2025,
            month=1,
            day=1,
            mindate=date(2025, 1, 1),
            maxdate=date(2030, 12, 31),
            firstweekday="sunday",
            weekendbackground="#E6F5EC",
            weekendforeground=NAVER_GREEN,
            headersbackground=NAVER_GREEN,
            headersforeground="white",
            selectbackground=NAVER_GREEN,
            selectforeground="white",
            background="white",
            bordercolor="#E1E4E7",
            othermonthforeground="#8C8D8D",
        )
        self.cal.pack(fill="both", expand=True, padx=8, pady=8)

        controls = ttk.Frame(frm, style="Card.TFrame")
        controls.grid(row=3, column=0, columnspan=6, sticky="ew")

        self.start_var = tk.StringVar(value="")
        self.end_var = tk.StringVar(value="")

        ttk.Button(
            controls,
            text="시작 날짜 지정",
            style="Secondary.TButton",
            command=self.pick_start,
        ).pack(side="left", padx=4, pady=4)
        ttk.Button(
            controls,
            text="종료 날짜 지정",
            style="Secondary.TButton",
            command=self.pick_end,
        ).pack(side="left", padx=4, pady=4)
        ttk.Label(
            controls,
            textvariable=self.start_var,
            style="Naver.TLabel",
        ).pack(side="left", padx=(16, 4))
        ttk.Label(
            controls,
            textvariable=self.end_var,
            style="Naver.TLabel",
        ).pack(side="left", padx=4)

        ttk.Button(
            frm,
            text="시험 기간 확정",
            style="Naver.TButton",
            command=self.set_period,
        ).grid(row=3, column=4, columnspan=2, sticky="e", padx=4)

        self.dates_container = ttk.Frame(frm, style="Card.TFrame")
        self.dates_container.grid(row=4, column=0, columnspan=6, sticky="nsew", pady=(16, 0))

    def pick_start(self):
        self.start_var.set(self.cal.selection_get().isoformat())

    def pick_end(self):
        self.end_var.set(self.cal.selection_get().isoformat())

    def set_period(self):
        try:
            start_date = datetime.strptime(self.start_var.get(), "%Y-%m-%d").date()
            end_date = datetime.strptime(self.end_var.get(), "%Y-%m-%d").date()
        except ValueError:
            messagebox.showerror("오류", "시작과 종료 날짜를 지정하세요")
            return
        if end_date < start_date:
            messagebox.showerror("오류", "종료일이 시작일보다 빠릅니다")
            return

        self.grade_count = int(self.grade_spin.get())
        self.class_count = int(self.class_spin.get())
        self.max_periods = int(self.max_period_spin.get())

        for child in self.dates_container.winfo_children():
            child.destroy()
        self.exam_dates.clear()

        ttk.Label(
            self.dates_container,
            text="날짜별 운영 교시",
            style="NaverHeading.TLabel",
        ).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 12))

        delta = (end_date - start_date).days
        row = 1
        for i in range(delta + 1):
            current = start_date + timedelta(days=i)
            if current.weekday() >= 5:
                continue
            date_str = current.isoformat()
            ttk.Label(
                self.dates_container,
                text=date_str,
                style="Naver.TLabel",
            ).grid(row=row, column=0, sticky="w", pady=4)
            spin = ttk.Spinbox(
                self.dates_container,
                from_=1,
                to=self.max_periods,
                width=6,
                style="Naver.TSpinbox",
                justify="center",
            )
            spin.set(str(self.max_periods))
            spin.grid(row=row, column=1, sticky="w", padx=(12, 0), pady=4)
            self.exam_dates[date_str] = spin
            row += 1

    # ----- 과목 -----
    def build_subject_frame(self):
        frm = self.subject_frame
        frm.columnconfigure(0, weight=1)
        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(3, weight=1)

        ttk.Label(frm, text="과목 등록", style="NaverHeading.TLabel").grid(
            row=0, column=0, columnspan=3, sticky="w", pady=(0, 12)
        )
        ttk.Label(frm, text="과목명", style="Naver.TLabel").grid(row=1, column=0, sticky="w")
        ttk.Label(frm, text="담당교사", style="Naver.TLabel").grid(row=1, column=1, sticky="w")

        self.subject_name = ttk.Entry(frm, style="Naver.TEntry")
        self.subject_teacher = ttk.Entry(frm, style="Naver.TEntry")
        self.subject_name.grid(row=2, column=0, sticky="ew", padx=(0, 12), pady=(0, 12))
        self.subject_teacher.grid(row=2, column=1, sticky="ew", padx=(0, 12), pady=(0, 12))

        ttk.Button(
            frm, text="과목 추가", style="Naver.TButton", command=self.add_subject
        ).grid(row=2, column=2, sticky="ew")

        list_card = ttk.Frame(frm, style="Card.TFrame")
        list_card.grid(row=3, column=0, columnspan=3, sticky="nsew", pady=10)

        self.subject_list = tk.Listbox(
            list_card,
            width=60,
            height=14,
            bg=NAVER_CARD,
            fg=NAVER_DARK,
            font=("Pretendard", 10),
            relief="flat",
            highlightthickness=0,
            selectbackground="#D7F2E2",
            activestyle="none",
        )
        self.subject_list.pack(fill="both", expand=True, padx=6, pady=6)

        action = ttk.Frame(frm, style="Card.TFrame")
        action.grid(row=4, column=0, columnspan=3, sticky="e")
        ttk.Button(action, text="저장", style="Secondary.TButton", command=self.save_subjects).pack(
            side="left", padx=4
        )
        ttk.Button(action, text="로드", style="Secondary.TButton", command=self.load_subjects).pack(
            side="left", padx=4
        )

    def add_subject(self):
        name = self.subject_name.get().strip()
        teacher = self.subject_teacher.get().strip()
        if not name or not teacher:
            messagebox.showwarning("경고", "과목명과 담당교사를 입력하세요")
            return
        self.subjects.append(Subject(name, teacher))
        self.subject_list.insert(tk.END, f"{name} - {teacher}")
        self.subject_name.delete(0, tk.END)
        self.subject_teacher.delete(0, tk.END)

    def save_subjects(self):
        if not self.subjects:
            messagebox.showwarning("경고", "저장할 과목이 없습니다")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".json", filetypes=[("JSON", "*.json")]
        )
        if not path:
            return
        payload = [{"name": s.name, "teacher": s.teacher} for s in self.subjects]
        with open(path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)

    def load_subjects(self):
        path = filedialog.askopenfilename(filetypes=[("JSON", "*.json")])
        if not path:
            return
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
        except Exception as exc:
            messagebox.showerror("오류", str(exc))
            return
        self.subjects.clear()
        self.subject_list.delete(0, tk.END)
        for item in data:
            subject = Subject(item["name"], item["teacher"])
            self.subjects.append(subject)
            self.subject_list.insert(tk.END, f"{subject.name} - {subject.teacher}")

    # ----- 교사 -----
    def build_teacher_frame(self):
        frm = self.teacher_frame
        frm.columnconfigure(0, weight=1)
        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(3, weight=1)

        ttk.Label(frm, text="교사 관리", style="NaverHeading.TLabel").grid(
            row=0, column=0, columnspan=3, sticky="w", pady=(0, 12)
        )
        ttk.Label(frm, text="교사 이름", style="Naver.TLabel").grid(row=1, column=0, sticky="w")

        self.teacher_name = ttk.Entry(frm, style="Naver.TEntry")
        self.teacher_name.grid(row=2, column=0, sticky="ew", padx=(0, 12), pady=(0, 12))

        ttk.Button(
            frm, text="교사 추가", style="Naver.TButton", command=self.add_teacher
        ).grid(row=2, column=1, sticky="ew")

        list_card = ttk.Frame(frm, style="Card.TFrame")
        list_card.grid(row=3, column=0, columnspan=3, sticky="nsew", pady=10)

        self.teacher_list = tk.Listbox(
            list_card,
            width=40,
            height=14,
            bg=NAVER_CARD,
            fg=NAVER_DARK,
            font=("Pretendard", 10),
            relief="flat",
            highlightthickness=0,
            selectbackground="#D7F2E2",
            activestyle="none",
        )
        self.teacher_list.pack(fill="both", expand=True, padx=6, pady=6)

        ttk.Button(
            frm,
            text="불참 시간 추가",
            style="Secondary.TButton",
            command=self.add_unavailability,
        ).grid(row=4, column=0, columnspan=3, sticky="ew", pady=4)

        action = ttk.Frame(frm, style="Card.TFrame")
        action.grid(row=5, column=0, columnspan=3, sticky="e")
        ttk.Button(action, text="저장", style="Secondary.TButton", command=self.save_teachers).pack(
            side="left", padx=4
        )
        ttk.Button(action, text="로드", style="Secondary.TButton", command=self.load_teachers).pack(
            side="left", padx=4
        )

    def add_teacher(self):
        name = self.teacher_name.get().strip()
        if not name:
            return
        if name in self.teachers:
            messagebox.showerror("오류", "이미 존재하는 교사")
            return
        self.teachers[name] = Teacher(name)
        self.teacher_list.insert(tk.END, name)
        self.teacher_name.delete(0, tk.END)

    def add_unavailability(self):
        selection = self.teacher_list.curselection()
        if not selection:
            messagebox.showwarning("경고", "교사를 선택하세요")
            return
        name = self.teacher_list.get(selection[0])
        teacher = self.teachers[name]

        win = tk.Toplevel(self)
        win.title(f"{name} 불가능 시간")
        win.configure(bg=NAVER_BG)

        ttk.Label(win, text="날짜 YYYY-MM-DD", style="Naver.TLabel").grid(
            row=0, column=0, padx=12, pady=(12, 6)
        )
        ttk.Label(win, text="교시", style="Naver.TLabel").grid(
            row=1, column=0, padx=12, pady=6
        )
        date_entry = ttk.Entry(win, width=18, style="Naver.TEntry")
        period_spin = ttk.Spinbox(
            win,
            from_=1,
            to=self.max_periods,
            width=6,
            style="Naver.TSpinbox",
            justify="center",
        )
        period_spin.set("1")
        date_entry.grid(row=0, column=1, padx=12, pady=(12, 6))
        period_spin.grid(row=1, column=1, padx=12, pady=6)

        def save():
            value = date_entry.get().strip()
            try:
                datetime.strptime(value, "%Y-%m-%d")
            except ValueError:
                messagebox.showerror("오류", "날짜 형식이 잘못되었습니다")
                return
            period = int(period_spin.get())
            teacher.add_unavailability(value, period)
            win.destroy()

        ttk.Button(win, text="저장", style="Naver.TButton", command=save).grid(
            row=2, column=0, columnspan=2, pady=(6, 12)
        )

    def save_teachers(self):
        if not self.teachers:
            messagebox.showwarning("경고", "저장할 교사가 없습니다")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".json", filetypes=[("JSON", "*.json")]
        )
        if not path:
            return
        payload = []
        for teacher in self.teachers.values():
            payload.append(
                {
                    "name": teacher.name,
                    "unavailable": [
                        {"date": d, "period": p} for d, p in teacher.unavailable
                    ],
                }
            )
        with open(path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)

    def load_teachers(self):
        path = filedialog.askopenfilename(filetypes=[("JSON", "*.json")])
        if not path:
            return
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
        except Exception as exc:
            messagebox.showerror("오류", str(exc))
            return
        self.teachers.clear()
        self.teacher_list.delete(0, tk.END)
        for item in data:
            teacher = Teacher(item["name"])
            for info in item.get("unavailable", []):
                teacher.add_unavailability(info["date"], info["period"])
            self.teachers[teacher.name] = teacher
            self.teacher_list.insert(tk.END, teacher.name)

    # ----- 스케줄 생성 -----
    def build_schedule_frame(self):
        frm = self.schedule_frame
        ttk.Label(frm, text="감독 스케줄", style="NaverHeading.TLabel").pack(
            anchor="w", pady=(0, 12)
        )
        buttons = ttk.Frame(frm, style="Card.TFrame")
        buttons.pack(fill="x")
        ttk.Button(
            buttons,
            text="스케줄 생성",
            style="Naver.TButton",
            command=self.generate_schedule,
        ).pack(side="left")
        ttk.Button(
            buttons,
            text="엑셀 저장",
            style="Secondary.TButton",
            command=self.export_excel,
        ).pack(side="left", padx=8)

        text_card = ttk.Frame(frm, style="Card.TFrame")
        text_card.pack(fill="both", expand=True, pady=(16, 0))
        self.schedule_text = tk.Text(
            text_card,
            width=90,
            height=26,
            bg=NAVER_CARD,
            fg=NAVER_DARK,
            font=("Pretendard", 10),
            relief="flat",
            highlightthickness=0,
        )
        self.schedule_text.pack(fill="both", expand=True, padx=6, pady=6)

    def generate_schedule(self):
        dates = []
        for date_str, spin in self.exam_dates.items():
            periods = int(spin.get())
            self.exam_dates[date_str] = periods
            for period in range(1, periods + 1):
                dates.append((date_str, period))

        if not dates:
            messagebox.showerror("오류", "시험 기간이 설정되지 않았습니다")
            return
        if not self.subjects:
            messagebox.showerror("오류", "과목이 없습니다")
            return
        if len(self.teachers) < 3:
            messagebox.showerror("오류", "교사가 최소 3명 필요합니다")
            return

        for teacher in self.teachers.values():
            teacher.assignments = {"main": 0, "assistant": 0, "selfstudy": 0}

        self.generated_schedule = []
        self.schedule_text.delete("1.0", tk.END)

        subject_iter = iter(self.subjects)
        for date_str, period in dates:
            try:
                subject = next(subject_iter)
                subject_name = subject.name
                forbidden = subject.teacher
            except StopIteration:
                subject = None
                subject_name = "자율학습"
                forbidden = None

            available = [
                teacher
                for teacher in self.teachers.values()
                if (date_str, period) not in teacher.unavailable and teacher.name != forbidden
            ]
            if len(available) < 3:
                messagebox.showerror(
                    "오류", f"{date_str} {period}교시에 배정 가능한 교사가 부족합니다"
                )
                self.generated_schedule = []
                return

            available.sort(
                key=lambda t: (
                    t.assignments["main"],
                    t.assignments["assistant"],
                    t.assignments["selfstudy"],
                    t.name,
                )
            )
            main, assistant, selfstudy = available[:3]
            main.assignments["main"] += 1
            assistant.assignments["assistant"] += 1
            selfstudy.assignments["selfstudy"] += 1

            self.generated_schedule.append(
                (date_str, period, subject_name, main.name, assistant.name, selfstudy.name)
            )

        self.render_schedule()

    def render_schedule(self):
        lines = []
        for item in self.generated_schedule:
            date_str, period, subject_name, main, assistant, selfstudy = item
            lines.append(
                f"{date_str} {period}교시 - {subject_name} : 정감독 {main}, 부감독 {assistant}, 자율학습 {selfstudy}"
            )
        lines.append("")
        lines.append("[교사별 배정 합계]")
        for teacher in self.teachers.values():
            totals = teacher.assignments
            lines.append(
                f"{teacher.name}: 정감독 {totals['main']}, 부감독 {totals['assistant']}, 자율학습 {totals['selfstudy']}"
            )
        self.schedule_text.delete("1.0", tk.END)
        self.schedule_text.insert(tk.END, "\n".join(lines))

    def export_excel(self):
        if not self.generated_schedule:
            messagebox.showwarning("경고", "스케줄이 없습니다")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")]
        )
        if not path:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "감독표"

        header_fill = PatternFill("solid", fgColor="03C75A")
        header_font = Font(color="FFFFFF", bold=True)
        border = Border(
            left=Side(style="thin", color="D0D3D4"),
            right=Side(style="thin", color="D0D3D4"),
            top=Side(style="thin", color="D0D3D4"),
            bottom=Side(style="thin", color="D0D3D4"),
        )

        ws.append(["날짜", "교시", "과목", "정감독", "부감독", "자율학습", "비고"])
        for row_index, row_data in enumerate(self.generated_schedule, start=2):
            ws.append(list(row_data) + [""])
            for col in range(1, 8):
                cell = ws.cell(row=row_index, column=col)
                cell.alignment = Alignment(vertical="center", horizontal="center")
                cell.border = border

        for col in range(1, 8):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        summary_start = 9
        summary_headers = ["교사", "정감독", "부감독", "자율학습"]
        for offset, text in enumerate(summary_headers):
            cell = ws.cell(row=1, column=summary_start + offset, value=text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        row_ptr = 2
        for teacher in self.teachers.values():
            ws.cell(row=row_ptr, column=summary_start, value=teacher.name)
            ws.cell(row=row_ptr, column=summary_start + 1, value=teacher.assignments["main"])
            ws.cell(row=row_ptr, column=summary_start + 2, value=teacher.assignments["assistant"])
            ws.cell(row=row_ptr, column=summary_start + 3, value=teacher.assignments["selfstudy"])
            for col in range(summary_start, summary_start + 4):
                cell = ws.cell(row=row_ptr, column=col)
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")
            row_ptr += 1

        info_row = row_ptr + 1
        ws.cell(row=info_row, column=summary_start, value="학년 수")
        ws.cell(row=info_row, column=summary_start + 1, value=self.grade_count)
        ws.cell(row=info_row + 1, column=summary_start, value="학급 수")
        ws.cell(row=info_row + 1, column=summary_start + 1, value=self.class_count)
        for row in (info_row, info_row + 1):
            for col in range(summary_start, summary_start + 2):
                cell = ws.cell(row=row, column=col)
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")

        for col in range(1, summary_start + 3):
            ws.column_dimensions[get_column_letter(col)].width = 16

        wb.save(path)


if __name__ == "__main__":
    app = ExamScheduler()
    app.mainloop()
